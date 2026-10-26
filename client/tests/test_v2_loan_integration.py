"""Loan raw use, principal, interest and cash reports remain separate and exact."""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from parsetrail.core.v2_cash_movement import read_cash_movement
from parsetrail.core.v2_profile_store import V2ProfileStore
from parsetrail.core.v2_report_drilldown import read_reviewed_entry_drilldown
from parsetrail.core.v2_report_export import write_reviewed_journal_workbook
from parsetrail.core.v2_source_pipeline import build_source_evidence

from devtools.parser_engine_pilot import test_wells_loan_lines as wells
from devtools.parser_engine_pilot.synthetic_family_sources import family_sources

from .test_v2_loan_payments import INTEREST, PREFERENCES, _post, _profile

JAN, FEB = date(2026, 1, 1), date(2026, 2, 1)


def _opening(store, account_id, amount):
    frame = store.account_anchor_snapshot(account_id)
    store.post_reviewed_manual_opening(
        account_id,
        amount_minor=amount,
        cutoff_date=date(2025, 12, 31),
        tracking_scope=frame.account.tracking_scope,
        reference="Synthetic reviewed opening",
        reason="Known opening",
        expected_decision_version=None,
        expected_basis_sha256=frame.basis_sha256,
    )


def _dict_rows(sheet):
    rows = iter(sheet.values)
    headers = next(rows)
    return [dict(zip(headers, row, strict=True)) for row in rows]


def test_cash_principal_and_interest_purposes_export_distinct_evidence_dimensions(tmp_path, app):
    profile = tmp_path / "loan.db"
    bank, rows = _profile(profile)
    with V2ProfileStore.open(profile) as store:
        _opening(store, "bank", 100000)
        _opening(store, "loan", -970000)
        event = _post(store, bank, rows, reference="=1+1")
        first = store.account_positions(as_of=date(2026, 1, 14))
        later = store.account_positions(as_of=date(2026, 1, 15))
        assert first.transfer_clearing_debit_minor == 22000 and later.transfer_clearing_debit_minor == 0
        assert {r.account_id: r.position_debit_minor for r in later.accounts} == {"bank": 78000, "loan": -950000}
        assert store.reviewed_category_summary(kind="expense", start=JAN, end=date(2026, 1, 15)).total_minor == 0
        assert store.reviewed_category_summary(kind="expense", start=date(2026, 1, 15), end=FEB).total_minor == 2000
        report = store.reviewed_journal_export(start=JAN, end=FEB)
        assert report.activity_rows_without_full_allocation == 0
        assert rows[0] in dict(report.source_rows)
        assert report.loan_source_links[0].derived_observation_id not in dict(report.source_rows)
    before = profile.read_bytes()
    cash = read_cash_movement(profile, ("bank",), start=JAN, end=FEB)
    assert cash.total_minor == -22000 and len(cash.rows) == 1
    assert cash.rows[0].classification == "loan_payment"
    assert {p.classification: p.amount_minor for p in cash.rows[0].components} == {
        "loan_principal": -20000,
        "expense": -2000,
    }
    assert sum(b.amount_minor for b in cash.buckets) == cash.total_minor
    assert "adds no expense posting" in cash.rows[0].event_note
    from parsetrail.gui.v2_cash_movement import V2CashMovementDialog
    from PySide6.QtCore import QDate, Qt

    dialog = V2CashMovementDialog(profile)
    dialog.start.setDate(QDate(2026, 1, 1))
    dialog.end.setDate(QDate(2026, 2, 1))
    dialog.accounts.item(0).setCheckState(Qt.Checked)
    dialog.entries.selectRow(0)
    app.processEvents()
    assert len(dialog.report.rows) == 1 and dialog.summary.rowCount() == 2
    assert all(dialog.summary.item(index, 1).text() == "1" for index in range(2))
    assert "loan payment" in dialog.scope.text() and "counts are not additive" in dialog.scope.text()
    assert dialog.sources.count() == 2
    dialog.close()
    assert read_cash_movement(profile, ("bank",), start=date(2026, 1, 15), end=FEB).total_minor == 0
    destination = tmp_path / "loan.xlsx"
    write_reviewed_journal_workbook(report, destination)
    workbook = load_workbook(destination, data_only=False)
    raw = _dict_rows(workbook["Source Allocations"])
    principal = _dict_rows(workbook["Derived Allocations"])
    uses = _dict_rows(workbook["Loan Source Uses"])
    assert len(raw) == len(principal) == len(uses) == 1
    assert raw[0]["Source row ID"] == bank and raw[0]["Signed minor units"] == "-22000"
    assert principal[0]["Parent source row ID"] == rows[0]
    assert principal[0]["Signed principal minor units"] == "20000"
    assert uses[0]["Raw signed payment minor units"] == "-22000"
    assert uses[0]["Principal debit minor units"] == "20000" and uses[0]["Interest expense minor units"] == "2000"
    assert uses[0]["Accounting date provenance"] == "assumed"
    assert workbook["Loan Source Locations"].max_row > 4
    assert sum(int(row["Net expense minor units"]) for row in _dict_rows(workbook["Reviewed Entries"])) == 2000
    reference_cell = workbook["Loan Events"].cell(2, 13)
    assert reference_cell.value == "=1+1" and reference_cell.data_type == "s"
    assert _dict_rows(workbook["Loan Events"])[0]["Event ID"] == event.event_id
    workbook.close()
    assert profile.read_bytes() == before


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


def test_interest_only_has_source_preview_without_fake_financial_observation_or_loan_position(tmp_path, app):
    profile = tmp_path / "interest.db"
    evidence = wells.parse()
    bank, rows = _profile(
        profile, amount=100, evidence=evidence, definition=wells.DEFINITION, raw=b"%PDF-synthetic-loan"
    )
    with V2ProfileStore.open(profile) as store:
        event = _post(store, bank, rows[:1])
        report = store.reviewed_journal_export(start=JAN, end=FEB)
        assert report.activity_rows_without_full_allocation == 1  # Unselected printed principal row remains pending.
        assert report.loan_source_links[0].derived_observation_id is None
        assert all(p.account_id != "loan" for entry in report.entries for p in entry.postings)
        assert (
            next(
                r for r in store.account_positions(as_of=date(2026, 1, 31)).accounts if r.account_id == "loan"
            ).position_debit_minor
            is None
        )
    before = profile.read_bytes()
    result = read_reviewed_entry_drilldown(profile, INTEREST, start=JAN, end=FEB)
    assert result.total_minor == 100 and len(result.entries) == 1
    assert not any(p.allocations for p in result.entries[0].postings)
    assert event.entry_keys[0] in {e.key for e in result.related_entries(result.entries[0])}
    from parsetrail.gui.v2_report_drilldown import V2ReviewedEntryDrilldown

    dialog = V2ReviewedEntryDrilldown(profile, INTEREST, start=JAN, end=FEB)
    app.processEvents()
    assert dialog.source.findData(evidence.source_sha256) >= 0 and dialog.preview.isEnabled()
    assert dialog.source.count() == 2
    assert "No principal financial observation/posting" in dialog.details.toPlainText()
    dialog.close()
    cash = read_cash_movement(profile, ("bank",), start=JAN, end=FEB)
    assert cash.total_minor == -100
    assert [(part.classification, part.amount_minor) for part in cash.rows[0].components] == [("expense", -100)]
    write_reviewed_journal_workbook(report, tmp_path / "interest.xlsx")
    workbook = load_workbook(tmp_path / "interest.xlsx")
    assert workbook["Derived Allocations"].max_row == 1
    assert len(_dict_rows(workbook["Loan Source Uses"])) == 1
    workbook.close()
    assert profile.read_bytes() == before


def test_whole_replacement_and_withdrawal_keep_audit_context_without_extra_money(tmp_path):
    profile = tmp_path / "history.db"
    bank, rows = _profile(profile)
    with V2ProfileStore.open(profile) as store:
        original = _post(store, bank, rows)
        basis = store.loan_review_snapshot(PREFERENCES, replacing_event_id=original.event_id).basis_sha256
        replacement = store.replace_reviewed_loan(
            original.event_id,
            bank,
            rows,
            expected_event_version=original.version,
            expected_basis_sha256=basis,
            preferences=PREFERENCES,
            reference="Corrected reviewed reference",
            reason="Replace entire payment",
        )
    result = read_reviewed_entry_drilldown(profile, INTEREST, start=JAN, end=FEB)
    assert result.total_minor == 2000 and len(result.entries) == 1
    related = {e.key for e in result.related_entries(result.entries[0])}
    assert set(original.entry_keys) <= related and replacement.entry_keys[0] in related
    assert len(result.report.loan_events) == 2 and len(result.report.loan_source_links) == 2
    with V2ProfileStore.open(profile) as store:
        store.withdraw_reviewed_loan(
            replacement.event_id, expected_event_version=replacement.version, reason="Withdraw entire replacement"
        )
        report = store.reviewed_journal_export(start=JAN, end=FEB)
        assert report.entries == () and report.activity_rows_without_full_allocation == 2
        assert all(not link.active for link in report.loan_source_links)
    assert read_cash_movement(profile, ("bank",), start=JAN, end=FEB).rows == ()
    write_reviewed_journal_workbook(report, tmp_path / "withdrawn.xlsx")
    workbook = load_workbook(tmp_path / "withdrawn.xlsx")
    assert workbook["Reviewed Entries"].max_row == workbook["Derived Allocations"].max_row == 1
    assert len(_dict_rows(workbook["Loan Source Uses"])) == 2
    workbook.close()


def test_actual_bounded_pdf_components_survive_posting_and_verified_restore(tmp_path):
    case = next(case for case in family_sources() if case.definition_name == "capital-one-loan-pilot-definition.json")
    definition = (Path(__file__).parents[2] / "devtools/parser_engine_pilot" / case.definition_name).read_bytes()
    evidence = build_source_evidence(case.source, definition, suffix=case.suffix)
    opening = evidence.payload.accounts[0].opening
    assert opening is None or opening.minor is None
    original_row = evidence.payload.accounts[0].rows[0]
    assert original_row.posting_date_guess is None and original_row.date_role == "activity_date"
    profile = tmp_path / "bounded-loan.db"
    bank, rows = _profile(profile, evidence=evidence, definition=definition, raw=case.source)
    with V2ProfileStore.open(profile) as store:
        event = _post(store, bank, rows)
        assert store.load_import(evidence.source_sha256).evidence == evidence
        assert event.manifest.components[0].row_evidence == original_row.model_dump_json()
    restored = tmp_path / "restored.db"
    V2ProfileStore.snapshot_to(profile, restored)
    with V2ProfileStore.open(restored, read_only=True) as store:
        assert store.load_import(evidence.source_sha256).evidence == evidence
        assert store.loan_history()[0] == event
        position = next(
            row for row in store.account_positions(as_of=date(2026, 1, 31)).accounts if row.account_id == "loan"
        )
        assert position.position_debit_minor is None  # Closing principal does not invent an opening.
        assert store.reviewed_category_summary(kind="expense", start=JAN, end=FEB).total_minor == 2000
