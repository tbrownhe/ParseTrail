"""Settlement cash parts preserve raw selected movement and source-dated fees."""

from datetime import date

from openpyxl import load_workbook
from parsetrail.core.v2_cash_movement import read_cash_movement
from parsetrail.core.v2_preferences import V2Preferences
from parsetrail.core.v2_profile_store import FinancialAccount, V2ProfileStore
from parsetrail.core.v2_report_export import write_reviewed_journal_workbook
from parsetrail.core.v2_settlements import SettlementFee
from parsetrail.core.v2_source_pipeline import build_source_evidence
from parsetrail.core.v2_table_card_source_output import from_anchored_card

from devtools.parser_engine_pilot.test_table_card_source_output import DEFINITION as CARD_DEFINITION
from devtools.parser_engine_pilot.test_table_card_source_output import SOURCE as CARD_SOURCE
from devtools.parser_engine_pilot.test_table_card_source_output import evidence as card_channels

from .test_v2_transfers import BANK_DEFINITION

JAN, FEB = date(2026, 1, 1), date(2026, 2, 1)


def _assets(profile, receipts=(6000, 3000)):
    source = (
        "Date,Account,Description,Amount,ReportedBalance\n"
        "01/01/2026,Checking,Outgoing,-$100.00,$900.00\n"
        + "".join(
            f"01/{index * 2 + 3:02d}/2026,Receiving {index},Receipt,${minor // 100}.{minor % 100:02d},$1000.00\n"
            for index, minor in enumerate(receipts)
        )
    ).encode()
    evidence = build_source_evidence(source, BANK_DEFINITION, suffix=".csv")
    with V2ProfileStore.create(profile) as store:
        store.import_source(source, BANK_DEFINITION, evidence, suffix=".csv")
        for identity, account_id in (
            ("Checking", "cash"),
            *((f"Receiving {index}", f"receiving:{index}") for index in range(len(receipts))),
        ):
            store.add_financial_account(FinancialAccount(account_id, identity, "asset", "synthetic"))
            store.bind_source_account(evidence.source_sha256, identity, "opaque", account_id)
        store.add_category_account("expense:fee", "Transfer fee", "expense")
    return tuple(row.row_id for row in evidence.payload.rows)


def _post(profile, row_ids, fees):
    with V2ProfileStore.open(profile) as store:
        frame = store.settlement_review_snapshot(preferences=V2Preferences())
        return store.post_reviewed_settlement(
            row_ids[0],
            tuple(row_ids[1:]),
            fees=fees,
            preferences=V2Preferences(),
            expected_basis_sha256=frame.basis_sha256,
            reason="Reviewed exact settlement",
        )


def _read(profile, accounts, start=JAN, end=FEB):
    return read_cash_movement(profile, tuple(accounts), start=start, end=end)


def _parts(report):
    return {bucket.classification: bucket.amount_minor for bucket in report.buckets}


def test_partial_outgoing_fee_replacement_keeps_single_effect_and_unresolved_clearing(tmp_path):
    profile = tmp_path / "partial.db"
    outgoing, first, second = _assets(profile)
    fees = (SettlementFee(outgoing, "expense:fee", 1000, "Reviewed sending fee", "Exact source allocation"),)
    original = _post(profile, (outgoing, first), fees)
    before = profile.read_bytes()
    cash = _read(profile, ("cash",))
    assert cash.total_minor == -10000 and len(cash.rows) == 1
    assert _parts(cash) == {"expense": -1000, "unsettled_transfer": -9000}
    assert sum(bucket.entry_count for bucket in cash.buckets) == 2
    assert cash.rows[0].classification == "mixed_settlement"
    assert cash.outside_scope_entries == 1
    both = _read(profile, ("cash", "receiving:0"))
    assert both.total_minor == -4000 and len(both.rows) == 2
    assert _parts(both) == {"expense": -1000, "unsettled_transfer": -3000}
    assert profile.read_bytes() == before
    with V2ProfileStore.open(profile) as store:
        position = store.account_positions(as_of=date(2026, 1, 31))
        assert position.transfer_clearing_debit_minor == 3000
        assert store.reviewed_category_summary(kind="expense", start=JAN, end=FEB).total_minor == 1000
        frame = store.settlement_review_snapshot(preferences=V2Preferences(), replacing_event_id=original)
        replacement = store.replace_reviewed_settlement(
            original,
            outgoing,
            (first, second),
            fees=fees,
            preferences=V2Preferences(),
            expected_basis_sha256=frame.basis_sha256,
            reason="Later receipt now evidenced",
        )
        assert store.account_positions(as_of=date(2026, 1, 31)).transfer_clearing_debit_minor == 0
        assert store.account_positions(as_of=date(2026, 1, 2)).transfer_clearing_debit_minor == 9000
    all_cash = _read(profile, ("cash", "receiving:0", "receiving:1"))
    assert all_cash.total_minor == -1000 and len(all_cash.rows) == 3
    assert _parts(all_cash) == {"expense": -1000, "within_scope_transfer": 0}
    assert {row.entry.event_id for row in all_cash.rows} == {replacement}
    assert all(sum(part.amount_minor for part in row.components) == row.amount_minor for row in all_cash.rows)
    backup = V2ProfileStore.snapshot_to(profile, tmp_path / "partial-backup.db")
    assert _read(backup, ("cash", "receiving:0", "receiving:1")) == all_cash


def test_receiving_asset_fee_decomposes_net_receipt_at_its_own_date_only(tmp_path):
    profile = tmp_path / "receiving.db"
    outgoing, incoming = _assets(profile, (9000,))
    _post(
        profile,
        (outgoing, incoming),
        (SettlementFee(incoming, "expense:fee", 1000, "Reviewed receiving deduction", "Exact net receipt allocation"),),
    )
    out_only = _read(profile, ("cash",))
    assert out_only.total_minor == -10000 and _parts(out_only) == {"owned_transfer": -10000}
    receipt = _read(profile, ("receiving:0",), date(2026, 1, 3), date(2026, 1, 4))
    assert receipt.total_minor == 9000 and len(receipt.rows) == 1
    assert _parts(receipt) == {"expense": -1000, "owned_transfer": 10000}
    both = _read(profile, ("cash", "receiving:0"))
    assert both.total_minor == -1000 and len(both.rows) == 2 and both.outside_scope_entries == 0
    assert _parts(both) == {"expense": -1000, "within_scope_transfer": 0}
    before_receipt = _read(profile, ("cash", "receiving:0"), JAN, date(2026, 1, 3))
    assert before_receipt.total_minor == -10000
    assert _parts(before_receipt) == {"within_scope_transfer": -10000}


def test_receiving_card_fee_does_not_manufacture_selected_cash_on_card_date(tmp_path):
    profile = tmp_path / "card.db"
    source = b"Date,Account,Description,Amount,ReportedBalance\n01/03/2026,Checking,Payment,-$30.00,$970.00\n"
    bank = build_source_evidence(source, BANK_DEFINITION, suffix=".csv")
    card = from_anchored_card(CARD_DEFINITION, *card_channels())
    with V2ProfileStore.create(profile) as store:
        store.import_source(source, BANK_DEFINITION, bank, suffix=".csv")
        store.import_source(CARD_SOURCE, CARD_DEFINITION, card, suffix=".pdf")
        store.add_financial_account(FinancialAccount("cash", "Checking", "asset", "synthetic"))
        store.add_financial_account(FinancialAccount("card", "Card", "liability", "citi"))
        store.add_category_account("expense:fee", "Payment fee", "expense")
        store.bind_source_account(bank.source_sha256, "Checking", "opaque", "cash")
        account = card.payload.accounts[0]
        store.bind_source_account(card.source_sha256, account.source_identity, account.source_identity_form, "card")
    incoming = account.rows[0].row_id
    _post(
        profile,
        (bank.payload.rows[0].row_id, incoming),
        (SettlementFee(incoming, "expense:fee", 1000, "Reviewed card fee", "Fee deducted from payment"),),
    )
    full = _read(profile, ("cash",))
    assert full.total_minor == -3000 and len(full.rows) == 1
    assert _parts(full) == {"liability_payment": -3000}
    card_day = _read(profile, ("cash",), date(2026, 1, 5), date(2026, 1, 6))
    assert card_day.total_minor == 0 and card_day.rows == () and card_day.outside_scope_entries == 1
    with V2ProfileStore.open(profile, read_only=True) as store:
        assert (
            store.reviewed_category_summary(kind="expense", start=date(2026, 1, 5), end=date(2026, 1, 6)).total_minor
            == 1000
        )


def test_export_keeps_full_fee_provenance_without_counting_out_of_period_or_withdrawn_money(tmp_path):
    profile = tmp_path / "export.db"
    outgoing, incoming = _assets(profile, (9000,))
    event = _post(
        profile,
        (outgoing, incoming),
        (SettlementFee(incoming, "expense:fee", 1000, "=1+1", "Explicit receiving deduction"),),
    )
    period_end = date(2026, 1, 2)
    with V2ProfileStore.open(profile, read_only=True) as store:
        report = store.reviewed_journal_export(start=JAN, end=period_end)
    assert len(report.entries) == 1 and len(report.settlements) == 1
    path = write_reviewed_journal_workbook(report, tmp_path / "settlements.xlsx")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames[-3:] == ["Settlement Events", "Settlement Legs", "Settlement Fees"]
        assert dict(list(workbook["Scope"].values)[1:])["Settlement metadata"].startswith("Full reviewed manifests")
        entries = list(workbook["Reviewed Entries"].values)[1:]
        assert len(entries) == 1 and sum(int(row[7]) for row in entries) == 0
        assert [row[4] for row in list(workbook["Settlement Events"].values)[1:]] == ["10000"]
        legs = list(workbook["Settlement Legs"].values)[1:]
        assert [row[10] for row in legs] == ["-10000", "9000"]
        assert legs[1][8] == "2026-01-03" and legs[1][11] == "10000"
        fee = list(workbook["Settlement Fees"].values)[1]
        assert fee[3] == "2026-01-03" and fee[6] == "1000" and fee[7] == "=1+1"
        assert workbook["Settlement Fees"]["H2"].data_type == "s"
        assert [row[5] for row in list(workbook["Source Allocations"].values)[1:]] == ["-10000"]
    finally:
        workbook.close()
    with V2ProfileStore.open(profile) as store:
        frame = store.settlement_review_snapshot(preferences=V2Preferences())
        store.withdraw_reviewed_transfer(event, expected_basis_sha256=frame.basis_sha256, reason="Withdraw settlement")
        retired = store.reviewed_journal_export(start=JAN, end=period_end)
    assert retired.entries == () and retired.settlements[0].status == "cancelled"
    path = write_reviewed_journal_workbook(retired, tmp_path / "withdrawn.xlsx")
    workbook = load_workbook(path, read_only=True)
    try:
        assert len(list(workbook["Reviewed Entries"].values)) == 1
        assert list(workbook["Settlement Events"].values)[1][1] == "cancelled"
        assert len(list(workbook["Settlement Fees"].values)) == 2
    finally:
        workbook.close()


def test_cash_dialog_shows_parts_without_adding_record_counts_or_movement(tmp_path, monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from parsetrail.gui.v2_cash_movement import V2CashMovementDialog
    from PySide6.QtCore import QDate, Qt
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    profile = tmp_path / "cash-dialog.db"
    outgoing, incoming = _assets(profile, (6000,))
    _post(
        profile,
        (outgoing, incoming),
        (SettlementFee(outgoing, "expense:fee", 1000, "Reviewed sending fee", "Exact fee allocation"),),
    )
    before = profile.read_bytes()
    dialog = V2CashMovementDialog(profile)
    try:
        dialog.start.setDate(QDate(2026, 1, 1))
        dialog.end.setDate(QDate(2026, 2, 1))
        for index in range(dialog.accounts.count()):
            if dialog.accounts.item(index).data(Qt.UserRole) == "cash":
                dialog.accounts.item(index).setCheckState(Qt.Checked)
        assert dialog.report.total_minor == -10000
        assert dialog.entries.rowCount() == 1 and dialog.summary.rowCount() == 2
        assert "counts are not additive" in dialog.scope.text()
        assert "-$100.00" in dialog.total.text()
        details = dialog.details.toPlainText()
        assert "Reviewed cash expense: -$10.00" in details
        assert "Owned transfer with unresolved clearing: -$90.00" in details
        assert "Remaining clearing after all reviewed receipts (any date): 3000" in details
        assert profile.read_bytes() == before
    finally:
        dialog.close()
        app.processEvents()


def test_workspace_exposes_whole_settlement_review_and_guards_import_in_progress(tmp_path, monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from parsetrail.gui import v2_settlement_review, v2_workspace_window
    from parsetrail.gui.v2_workspace_window import V2WorkspaceWindow
    from PySide6.QtWidgets import QApplication

    from .test_v2_workspace_window import _workspace

    app = QApplication.instance() or QApplication([])
    workspace = _workspace(tmp_path)
    opened = []
    transfer_opens = []

    class Review:
        def __init__(self, selected_workspace, event_id=None, parent=None):
            opened.append((selected_workspace, event_id, parent))

        def exec(self):
            return 0

    monkeypatch.setattr(v2_settlement_review, "V2SettlementReview", Review)
    monkeypatch.setattr(v2_workspace_window.V2TransferReview, "exec", lambda dialog: transfer_opens.append(dialog) or 0)
    window = V2WorkspaceWindow(workspace)
    try:
        window.settlement_review_action.trigger()
        assert opened == [(workspace, None, window)]
        window.transfer_review_action.trigger()
        assert len(transfer_opens) == 1
        window.import_worker = object()
        window.refresh()
        assert not window.settlement_review_action.isEnabled()
        assert not window.transfer_review_action.isEnabled()
        assert not window.transfer_review_button.isEnabled()
        window._review_settlements()
        window._review_transfers()
        window.transfer_review_action.trigger()
        window.transfer_review_button.click()
        assert len(opened) == 1
        assert len(transfer_opens) == 1
    finally:
        window.import_worker = None
        window.close()
        app.processEvents()
