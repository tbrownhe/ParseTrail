"""Whole liability-register repayments are neither principal estimates nor expenses."""

from datetime import date

from openpyxl import load_workbook
from parsetrail.core.v2_cash_movement import read_cash_movement
from parsetrail.core.v2_preferences import V2Preferences
from parsetrail.core.v2_profile_store import V2ProfileStore
from parsetrail.core.v2_report_drilldown import read_reviewed_entry_drilldown
from parsetrail.core.v2_report_export import write_reviewed_journal_workbook

from .test_v2_liability_register import _post, _profile
from .test_v2_liability_register import yamaha_source as yamaha_source

FEB, MAR = date(2026, 2, 1), date(2026, 3, 1)


def _review_opening(store, source):
    store.record_statement_source_review(
        "register",
        source,
        expected_revision_id=None,
        opening_origin="reported",
        closing_origin="reported",
        timing_confirmed=True,
        posting_dates="reported",
        source_reference="Printed register endpoints and date columns",
    )
    store.post_reviewed_opening("register")


def test_register_expense_and_repayment_keep_cash_spending_and_whole_liability_distinct(tmp_path, yamaha_source):
    profile = tmp_path / "register.db"
    evidence, bank = _profile(profile, yamaha_source)
    charge, payment = evidence.payload.accounts[0].rows
    with V2ProfileStore.open(profile) as store:
        _review_opening(store, evidence.source_sha256)
        _post(store, evidence.source_sha256, charge.row_id)
        event_id = store.post_reviewed_transfer(
            bank, payment.row_id, preferences=V2Preferences(), reason="Confirmed owned repayment"
        )
        event = next(event for event in store.transfer_history() if event.event_id == event_id)
        assert event.pair.kind == "liability payment"
        frame = store.source_dated_value_snapshot("register")
        candidate = frame.source_value_candidates[0]
        store.record_reviewed_source_value(
            "register",
            candidate.candidate_id,
            expected_candidate_sha256=candidate.version,
            expected_basis_sha256=frame.basis_sha256,
            reference="Printed full liability closing",
            reason="Nonposting source value",
        )
        for day, expected, clearing in ((7, -101000, 0), (9, -101000, 2000), (10, -99000, 0)):
            positions = store.account_positions(as_of=date(2026, 2, day))
            register = next(row for row in positions.accounts if row.account_id == "register")
            assert register.position_debit_minor == expected and register.tracking_scope == "account_total"
            assert positions.transfer_clearing_debit_minor == clearing
        report = store.reviewed_journal_export(start=FEB, end=MAR)
        assert store.reviewed_category_summary(kind="expense", start=FEB, end=MAR).total_minor == 1000
        assert report.loan_events == () and report.loan_source_links == ()
        assert not store.connection.execute("SELECT 1 FROM Accounts WHERE id='builtin:loan-interest'").fetchone()
        assert not store.connection.execute("SELECT 1 FROM DerivedFinancialObservations").fetchone()
    before = profile.read_bytes()
    cash = read_cash_movement(profile, ("bank",), start=FEB, end=MAR)
    assert cash.total_minor == -2000 and len(cash.rows) == 1
    assert [(part.classification, part.amount_minor) for part in cash.rows[0].components] == [
        ("liability_payment", -2000)
    ]
    assert "no principal/interest split is inferred" in cash.rows[0].event_note
    assert "2026-02-09" in cash.rows[0].event_note and "2026-02-10" in cash.rows[0].event_note
    assert read_cash_movement(profile, ("bank",), start=date(2026, 2, 10), end=MAR).rows == ()
    expense = read_reviewed_entry_drilldown(profile, "expense:first", start=FEB, end=MAR)
    assert expense.total_minor == 1000 and len(expense.entries) == 1
    assert expense.entries[0].posting_date == date(2026, 2, 7)
    write_reviewed_journal_workbook(report, tmp_path / "register.xlsx")
    workbook = load_workbook(tmp_path / "register.xlsx", data_only=False)
    assert "Derived Allocations" not in workbook.sheetnames and "Loan Events" not in workbook.sheetnames
    assert "Settlement Events" not in workbook.sheetnames  # Existing whole-pair exports use raw allocation provenance.
    assert any(payment.row_id in row and "2000" in row for row in list(workbook["Source Allocations"].values)[1:])
    entries = list(workbook["Reviewed Entries"].values)
    expense_column = entries[0].index("Net expense minor units")
    assert sum(int(row[expense_column]) for row in entries[1:]) == 1000
    assert workbook["Source Allocations"].max_row == 4  # Charge, observed cash outflow and whole liability receipt.
    workbook.close()
    assert profile.read_bytes() == before


def test_generic_register_settlement_withdrawal_and_restore_preserve_zero_extra_expense(tmp_path, yamaha_source):
    profile = tmp_path / "settlement.db"
    evidence, bank = _profile(profile, yamaha_source)
    charge, payment = evidence.payload.accounts[0].rows
    with V2ProfileStore.open(profile) as store:
        _post(store, evidence.source_sha256, charge.row_id)
        frame = store.settlement_review_snapshot(preferences=V2Preferences())
        event_id = store.post_reviewed_settlement(
            bank,
            (payment.row_id,),
            preferences=V2Preferences(),
            expected_basis_sha256=frame.basis_sha256,
            reason="Reviewed observed register repayment",
        )
        event = next(event for event in store.transfer_history() if event.event_id == event_id)
        assert event.settlement.kind == "liability payment"
        assert store.reviewed_category_summary(kind="expense", start=FEB, end=MAR).total_minor == 1000
        store.withdraw_reviewed_transfer(
            event.event_id,
            expected_basis_sha256=store.settlement_review_snapshot(preferences=V2Preferences()).basis_sha256,
            reason="Withdraw entire repayment",
        )
        assert store._active_consumed() == {charge.row_id: -1000}
        fresh = store.post_reviewed_transfer(
            bank, payment.row_id, preferences=V2Preferences(), reason="Fresh whole repayment"
        )
        assert fresh != event.event_id
    restored = tmp_path / "restored.db"
    V2ProfileStore.snapshot_to(profile, restored)
    cash = read_cash_movement(restored, ("bank",), start=FEB, end=MAR)
    assert cash.total_minor == -2000 and len(cash.rows) == 1 and cash.rows[0].classification == "liability_payment"
    with V2ProfileStore.open(restored, read_only=True) as store:
        report = store.reviewed_journal_export(start=FEB, end=MAR)
        assert len(report.entries) == 3 and len(report.settlements) == 1
        assert (
            report.settlements[0].status == "cancelled"
        )  # Advanced history remains metadata; the fresh pair is active.
        assert len(store.transfer_history()) == 2
        assert store.reviewed_category_summary(kind="expense", start=FEB, end=MAR).total_minor == 1000
        assert store.load_import(evidence.source_sha256).evidence == evidence
