"""Cash-component books and expected destinations never fabricate receipts."""

from datetime import date

from openpyxl import load_workbook
from parsetrail.core.v2_cash_movement import read_cash_movement
from parsetrail.core.v2_preferences import V2Preferences
from parsetrail.core.v2_profile_store import V2ProfileStore
from parsetrail.core.v2_report_export import write_reviewed_journal_workbook
from parsetrail.core.v2_settlements import ExpectedOwnedDestination

from .test_v2_hsa_cash import hsa_profile, post_missing, source_review
from .test_v2_hsa_cash import hsa_source as hsa_source
from .test_v2_settlements import profile as bank_profile

AUG, SEP = date(2026, 8, 1), date(2026, 9, 1)
OCT, NOV = date(2026, 10, 1), date(2026, 11, 1)
PREFS = V2Preferences()


def _read(profile, accounts, start=AUG, end=SEP):
    return read_cash_movement(profile, tuple(accounts), start=start, end=end)


def test_full_hsa_cash_component_matches_own_endpoint_with_unreceived_investment_separate(tmp_path, hsa_source):
    profile = tmp_path / "hsa.db"
    evidence = hsa_profile(profile, hsa_source)
    rows = evidence.payload.accounts[0].rows
    with V2ProfileStore.open(profile) as store:
        source_review(store, evidence)
        opening = store.post_reviewed_opening("hsa")
        assert opening.tracking_scope == "cash_component" and opening.amount_minor == 100000
        for row in (rows[0], rows[2]):
            store.post_reviewed_cash_category(
                row.row_id,
                "income:first",
                purpose="income",
                reason="Confirmed receipt purpose, no inferred description",
            )
        event = post_missing(
            store,
            rows[1].row_id,
            destination=ExpectedOwnedDestination(
                "investment",
                "Reviewed owned investment destination; receipt unavailable",
                True,
            ),
        )
        frame = store.source_dated_value_snapshot("hsa")
        (candidate,) = frame.source_value_candidates
        store.record_reviewed_source_value(
            "hsa",
            candidate.candidate_id,
            expected_candidate_sha256=candidate.version,
            expected_basis_sha256=frame.basis_sha256,
            reference="Printed cash-component closing only",
            reason="Nonposting point independent of portfolio value",
        )
        (check,) = store.statement_reconciliation_checks()
        assert check.reconciled and check.balance_scope == check.tracking_scope == "cash_component"
        assert check.journal_closing_minor == check.printed_closing_minor == 110010
        store.save_reconciliation_decision(
            "hsa",
            evidence.source_sha256,
            expected_input_version=check.input_version,
            reason="Cash component only; missing receipt stays unresolved",
        )
        positions = store.account_positions(as_of=date(2026, 8, 31))
        hsa = next(row for row in positions.accounts if row.account_id == "hsa")
        destination = next(row for row in positions.accounts if row.account_id == "investment")
        assert hsa.recorded_debit_minor == hsa.position_debit_minor == hsa.latest_dated_value.amount_minor == 110010
        assert destination.recorded_debit_minor == 0 and destination.position_debit_minor is None
        assert positions.transfer_clearing_debit_minor == 10000
        assert store._active_consumed() == {row.row_id: row.amount.minor for row in rows}
        assert all(p.account_id != "investment" for entry in store._active_journal_entries() for p in entry.postings)
        assert store.reviewed_category_summary(kind="income", start=AUG, end=SEP).total_minor == 20010
        assert store.reviewed_category_summary(kind="expense", start=AUG, end=SEP).total_minor == 0
        assert rows[2].description.text is None
        assert store.load_import(evidence.source_sha256).evidence == evidence
        assert not evidence.payload.capture_verified
        assert store.financial_account_scope_locked("investment")
        assert store.transfer_history()[0].event_id == event
    before = profile.read_bytes()
    cash = _read(profile, ("hsa",))
    assert cash.total_minor == 10010 and len(cash.rows) == 3
    assert {bucket.classification: bucket.amount_minor for bucket in cash.buckets} == {
        "income": 20010,
        "unsettled_transfer": -10000,
    }
    assert cash.expected_receipts_without_source == 1
    receiver_only = _read(profile, ("investment",))
    assert receiver_only.total_minor == 0 and receiver_only.rows == ()
    assert receiver_only.expected_receipts_without_source == 1
    assert "no destination receipt or positive financial posting" in next(
        row.event_note for row in cash.rows if row.classification == "unsettled_transfer"
    )
    assert profile.read_bytes() == before
    backup = V2ProfileStore.snapshot_to(profile, tmp_path / "hsa-backup.db")
    assert _read(backup, ("hsa",)) == cash
    with V2ProfileStore.open(backup, read_only=True) as store:
        assert store.account_positions(as_of=date(2026, 8, 31)) == positions
        assert store.reconciliation_history("hsa")[0].check.tracking_scope == "cash_component"


def test_expected_destination_export_contains_metadata_and_one_actual_source_leg_only(tmp_path):
    profile = tmp_path / "expectation.db"
    evidence, outgoing, _incoming = bank_profile(profile)
    with V2ProfileStore.open(profile) as store:
        frame = store.settlement_review_snapshot(preferences=PREFS)
        event = store.post_reviewed_settlement(
            outgoing,
            (),
            preferences=PREFS,
            expected_basis_sha256=frame.basis_sha256,
            expected_destination=ExpectedOwnedDestination("savings", "=1+1", True),
            reason="Owned destination reviewed; no receipt evidence",
        )
        report = store.reviewed_journal_export(start=OCT, end=NOV)
        assert len(report.entries) == 1 and len(report.entries[0].postings) == 2
        assert {p.account_id for p in report.entries[0].postings} == {"checking", "internal:transfer-clearing:USD"}
        assert store._active_consumed() == {outgoing: -10000}
        assert store.load_import(evidence.source_sha256).evidence == evidence
    path = write_reviewed_journal_workbook(report, tmp_path / "expectation.xlsx")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        events = list(workbook["Settlement Events"].values)
        row = dict(zip(events[0], events[1], strict=True))
        assert row["Event ID"] == event and row["Received gross minor units"] == "0"
        assert row["Expected owned destination ID"] == "savings" and row["Ownership acknowledged"] is True
        assert row["Ownership review reference"] == "=1+1"
        column = events[0].index("Ownership review reference") + 1
        assert workbook["Settlement Events"].cell(2, column).data_type == "s"
        assert "no receipt evidence or destination posting" in row["Expectation qualification"]
        assert len(list(workbook["Settlement Legs"].values)) == 2
        assert list(workbook["Settlement Legs"].values)[1][2] == "outgoing"
        assert len(list(workbook["Source Allocations"].values)) == 2
        assert len(list(workbook["Posting Lines"].values)) == 3
    finally:
        workbook.close()


def test_actual_receipt_replacement_removes_current_expectation_without_double_counting_outflow(tmp_path):
    profile = tmp_path / "replacement.db"
    _evidence, outgoing, incoming = bank_profile(profile, receipts=(("10/05/2026", "10000"),))
    with V2ProfileStore.open(profile) as store:
        original = post_missing(store, outgoing)
    prior = _read(profile, ("checking", "savings"), OCT, NOV)
    assert prior.total_minor == -10000 and prior.expected_receipts_without_source == 1
    assert len(prior.rows) == 1
    with V2ProfileStore.open(profile) as store:
        frame = store.settlement_review_snapshot(preferences=PREFS, replacing_event_id=original)
        replacement = store.replace_reviewed_settlement(
            original,
            outgoing,
            incoming,
            preferences=PREFS,
            expected_basis_sha256=frame.basis_sha256,
            reason="Actual receipt now reviewed, no expected receipt substituted",
        )
        history = {event.event_id: event for event in store.transfer_history()}
        assert history[original].settlement.expected_destination.account_id == "savings"
        assert history[original].status == "replaced" and history[replacement].settlement.expected_destination is None
        assert store.account_positions(as_of=date(2026, 10, 4)).transfer_clearing_debit_minor == 10000
        assert store.account_positions(as_of=date(2026, 10, 5)).transfer_clearing_debit_minor == 0
    complete = _read(profile, ("checking", "savings"), OCT, NOV)
    assert complete.total_minor == 0 and complete.expected_receipts_without_source == 0
    assert len(complete.rows) == 2 and {row.entry.event_id for row in complete.rows} == {replacement}
    assert {bucket.classification: bucket.amount_minor for bucket in complete.buckets} == {"within_scope_transfer": 0}
    interim = _read(profile, ("checking", "savings"), OCT, date(2026, 10, 5))
    assert interim.total_minor == -10000 and len(interim.rows) == 1
