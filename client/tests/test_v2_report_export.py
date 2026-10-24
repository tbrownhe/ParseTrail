"""Reviewed journal exports retain active rows, exact postings and scope."""

from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from parsetrail.core.v2_profile_store import FinancialAccount, V2ProfileStore
from parsetrail.core.v2_report_export import ReportExportError, write_reviewed_journal_workbook
from parsetrail.core.v2_source_pipeline import build_source_evidence

DEFINITION = (
    Path(__file__).parents[2] / "devtools" / "v2_fresh_slice" / "synthetic-posted-cash-definition.json"
).read_bytes()
SOURCE = (
    b"Date,Account,Description,Amount,ReportedBalance\n"
    b"10/01/2026,Checking A,Coffee,-$4.00,$996.00\n"
    b"10/02/2026,Checking A,Deposit,$10.00,$1006.00\n"
    b"10/03/2026,Checking A,Tea,-$3.00,$1003.00\n"
)


def test_workbook_has_active_reviewed_entries_postings_corrections_and_scope(tmp_path: Path) -> None:
    evidence = build_source_evidence(SOURCE, DEFINITION, suffix=".csv")
    coffee, deposit, _tea = evidence.payload.rows
    profile = tmp_path / "fresh.db"
    with V2ProfileStore.create(profile) as store:
        store.import_source(SOURCE, DEFINITION, evidence, suffix=".csv")
        store.add_financial_account(FinancialAccount("cash", "Checking", "asset", "synthetic"))
        store.add_category_account("expense:food", "Food", "expense")
        store.add_category_account("expense:travel", "Travel", "expense")
        store.add_category_account("income:pay", "Pay", "income")
        store.bind_source_account(evidence.source_sha256, "Checking A", "opaque", "cash")
        original = store.post_reviewed_cash_category(coffee.row_id, "expense:food", reason="Reviewed coffee")
        replacement = store.correct_reviewed_cash_category(original, "expense:travel", reason="Correct category")
        store.post_reviewed_cash_category(deposit.row_id, "income:pay", reason="Confirmed income")
        store.rename_category_account("expense:travel", "=1+1")
    before = hashlib.sha256(profile.read_bytes()).hexdigest()
    with V2ProfileStore.open(profile, read_only=True) as store:
        export = store.reviewed_journal_export(start=date(2026, 10, 1), end=date(2026, 11, 1))
        assert {entry.key for entry in export.entries} == {replacement, f"review:cash:income:{deposit.row_id}"}
        assert export.activity_rows_without_full_allocation == 1
        assert export.sources_with_incomplete_activity == 1
        assert len(export.corrections) == 1
        expense = store.reviewed_category_summary(kind="expense", start=export.start, end=export.end)
        income = store.reviewed_category_summary(kind="income", start=export.start, end=export.end)
    path = write_reviewed_journal_workbook(export, tmp_path / "reviewed.xlsx")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Scope",
            "Reviewed Entries",
            "Posting Lines",
            "Source Allocations",
            "Corrections",
        ]
        entry_rows = list(workbook["Reviewed Entries"].values)
        assert len(entry_rows) == 3
        assert sum(row[6] for row in entry_rows[1:]) == expense.total_minor / 100
        assert sum(row[8] for row in entry_rows[1:]) == income.total_minor / 100
        posting_rows = list(workbook["Posting Lines"].values)
        assert len(posting_rows) == 5
        assert sum(int(row[7]) for row in posting_rows[1:] if row[5] == "expense") == expense.total_minor
        assert -sum(int(row[7]) for row in posting_rows[1:] if row[5] == "income") == income.total_minor
        assert any(row[4] == "=1+1" for row in posting_rows[1:])
        formula_name = next(cell for cell in workbook["Posting Lines"]["E"] if cell.value == "=1+1")
        assert formula_name.data_type == "s"
        allocation_rows = list(workbook["Source Allocations"].values)
        assert len(allocation_rows) == 3
        assert {row[3] for row in allocation_rows[1:]} == {evidence.source_sha256}
        scope = dict(list(workbook["Scope"].values)[1:])
        assert scope["Activity source rows without full allocation, whole profile"] == 1
        assert scope["Active reviewed entries"] == 2
        assert scope["Correction presentation"].startswith("Superseded")
        assert len(list(workbook["Corrections"].values)) == 2
    finally:
        workbook.close()
    assert hashlib.sha256(profile.read_bytes()).hexdigest() == before


def test_workbook_requires_new_xlsx_destination(tmp_path: Path) -> None:
    with V2ProfileStore.create(tmp_path / "fresh.db") as store:
        report = store.reviewed_journal_export(start=date(2026, 10, 1), end=date(2026, 11, 1))
    empty = write_reviewed_journal_workbook(report, tmp_path / "empty.xlsx")
    workbook = load_workbook(empty, read_only=True)
    try:
        assert len(list(workbook["Reviewed Entries"].values)) == 1
        assert dict(list(workbook["Scope"].values)[1:])["Active reviewed entries"] == 0
    finally:
        workbook.close()
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"keep")
    with pytest.raises(ReportExportError, match="report_destination_exists"):
        write_reviewed_journal_workbook(report, existing)
    assert existing.read_bytes() == b"keep"
    with pytest.raises(ReportExportError, match="report_destination_requires_xlsx"):
        write_reviewed_journal_workbook(report, tmp_path / "wrong.csv")
