"""Spreadsheet presentation for a reviewed 2.0 journal snapshot."""

from __future__ import annotations

import os
import tempfile
from itertools import chain
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from parsetrail.core.source_evidence import ObservationEvidence
from parsetrail.core.v2_profile_store import ReviewedJournalExport


class ReportExportError(ValueError):
    """The reviewed report cannot be published at the chosen path."""


def _sheet(workbook: Workbook, name: str, headers: tuple[str, ...], rows: list[tuple]) -> None:
    if len(rows) > 1_048_575:
        raise ReportExportError("report_sheet_row_limit")
    sheet = workbook.create_sheet(name)
    for row_number, values in enumerate(chain((headers,), rows), start=1):
        for column_number, value in enumerate(values, start=1):
            if isinstance(value, str) and len(value) > 32_767:
                raise ReportExportError("report_cell_too_long")
            cell = sheet.cell(row_number, column_number, value)
            if isinstance(value, str):
                cell.data_type = "s"  # Source descriptions and authored names must not become Excel formulas.
            if row_number == 1:
                cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[letter].width = min(max(width + 2, 12), 64)


def _loan_metadata(workbook: Workbook, report: ReviewedJournalExport, derived_allocations: list[tuple]) -> list[str]:
    """Full raw-row use and derived principal support are distinct audit dimensions."""
    if not report.loan_events:
        return []
    workbook["Scope"].append(
        (
            "Loan metadata",
            "Loan event/source-use/component history covers all dates, including retired events. It adds no money "
            "to the period's entries. Source Allocations retains raw financial rows; Derived Allocations identifies "
            "canonical principal effects. Whole raw payment use must not be summed as principal allocation. "
            "Lender accounting dates are explicitly reviewed assumptions based on retained activity dates.",
        )
    )
    _sheet(
        workbook,
        "Derived Allocations",
        (
            "Entry key",
            "Posting line",
            "Derived observation ID",
            "Parent source row ID",
            "Source SHA-256",
            "Signed principal USD",
            "Signed principal minor units",
            "Normalization rule",
        ),
        derived_allocations,
    )
    _sheet(
        workbook,
        "Loan Events",
        (
            "Event ID",
            "State",
            "Bank account ID",
            "Bank source row ID",
            "Bank source SHA-256",
            "Bank date",
            "Loan account ID",
            "Captured loan name",
            "Tracking scope",
            "Gross payment minor units",
            "Principal debit minor units",
            "Interest expense minor units",
            "Reference",
            "Reason",
            "Entry keys",
            "Replacement event",
            "Created at",
            "Version",
        ),
        [
            (
                event.event_id,
                event.status,
                event.manifest.bank.account_id,
                event.manifest.bank.row_id,
                event.manifest.bank.source_sha256,
                event.manifest.bank.posting_date,
                event.manifest.components[0].account_id,
                event.manifest.components[0].account_name,
                event.manifest.components[0].tracking_scope,
                str(-event.manifest.bank.amount_minor),
                str(event.manifest.principal_minor),
                str(event.manifest.interest_minor),
                event.reference,
                event.reason,
                "\n".join(event.entry_keys),
                event.replacement_event_id,
                event.created_at,
                event.version,
            )
            for event in report.loan_events
        ],
    )
    _sheet(
        workbook,
        "Loan Source Uses",
        (
            "Event ID",
            "Entry key",
            "Raw source row ID",
            "Source SHA-256",
            "Loan account ID",
            "Active use",
            "Raw signed payment minor units",
            "Derived principal observation ID",
            "Principal debit minor units",
            "Interest expense minor units",
            "Source activity date",
            "Accounting date",
            "Accounting date provenance",
            "Parser ID",
            "Sign factor",
            "Tracking scope",
        ),
        [
            (
                link.event_id,
                link.entry_key,
                link.row_id,
                link.source_sha256,
                link.account_id,
                link.active,
                str(link.raw_amount_minor),
                link.derived_observation_id,
                str(link.principal_minor),
                str(link.interest_minor),
                link.activity_date,
                link.effective_date,
                link.date_provenance,
                link.proof.parser_id,
                link.proof.sign_factor,
                link.proof.tracking_scope,
            )
            for link in report.loan_source_links
        ],
    )
    component_rows, locations = [], []
    for link in report.loan_source_links:
        row = ObservationEvidence.model_validate_json(link.proof.row_evidence)
        fields = (("raw payment", row.amount), *((part.role, part.amount) for part in row.components))
        for field, money in fields:
            component_rows.append(
                (link.event_id, link.row_id, field, str(money.minor), money.raw, money.provenance, "USD")
            )
        for field, value in (*fields, ("activity date", row.date), ("description", row.description)):
            for ordinal, location in enumerate(value.locations, start=1):
                locations.append((link.event_id, link.row_id, field, ordinal, location.model_dump_json()))
    _sheet(
        workbook,
        "Loan Components",
        ("Event ID", "Source row ID", "Field", "Raw signed minor units", "Printed text", "Provenance", "Currency"),
        component_rows,
    )
    _sheet(
        workbook,
        "Loan Source Locations",
        ("Event ID", "Source row ID", "Field", "Location ordinal", "Exact location JSON"),
        locations,
    )
    _sheet(
        workbook,
        "Loan Actions",
        ("Action ID", "Event ID", "Expected event version", "Replacement event", "Reason", "Created at"),
        [
            (
                row.action_id,
                row.event_id,
                row.expected_event_version,
                row.replacement_event_id,
                row.reason,
                row.created_at,
            )
            for row in report.loan_actions
        ],
    )
    return [
        "Derived Allocations",
        "Loan Events",
        "Loan Source Uses",
        "Loan Components",
        "Loan Source Locations",
        "Loan Actions",
    ]


def _value_metadata(workbook: Workbook, report: ReviewedJournalExport, accounts: dict) -> list[str]:
    """Nonposting points and their full source/audit context add no entry totals."""
    if not report.dated_values:
        return []
    workbook["Scope"].append(
        (
            "Dated value metadata",
            "Nonposting points filtered by value date, including superseded/withdrawn history. Creation/withdrawal "
            "dates do not add money. Current points can conflict; values and total/component scopes are never summed "
            "into entries, balances or certified net worth.",
        )
    )
    _sheet(
        workbook,
        "Dated Values",
        (
            "Value ID",
            "Account ID",
            "Account name",
            "Value date",
            "Origin",
            "Tracking scope",
            "Reviewed debit minor units",
            "State",
            "Reference",
            "Reason",
            "Corrects value",
            "Superseded by",
            "Withdrawal ID",
            "Created at",
            "Version",
            "Conflicting value acknowledged",
            "Source review changed since recording",
        ),
        [
            (
                value.value_id,
                value.account_id,
                accounts[value.account_id].name,
                value.value_date.isoformat(),
                value.origin,
                value.tracking_scope,
                str(value.amount_minor),
                value.state,
                value.reference,
                value.reason,
                value.corrects_value_id,
                value.superseded_by_value_id,
                value.withdrawal.withdrawal_id if value.withdrawal else None,
                value.created_at,
                value.version,
                value.allow_conflict,
                value.source_review_changed,
            )
            for value in report.dated_values
        ],
    )
    _sheet(
        workbook,
        "Dated Value Withdrawals",
        ("Withdrawal ID", "Value ID", "Account ID", "Prior value version", "Reason", "Created at", "Version"),
        [
            (
                row.withdrawal_id,
                row.value_id,
                row.account_id,
                row.prior_version,
                row.reason,
                row.created_at,
                row.version,
            )
            for row in report.dated_value_withdrawals
        ],
    )
    names = ["Dated Values", "Dated Value Withdrawals"]
    source_values = tuple(value for value in report.dated_values if value.source_proof is not None)
    if source_values:
        _sheet(
            workbook,
            "Source Value Proofs",
            (
                "Value ID",
                "Source SHA-256",
                "Evidence SHA-256",
                "Definition SHA-256",
                "Parser ID",
                "Definition revision",
                "Source section identity",
                "Identity form",
                "Financial kind",
                "Tracking scope",
                "Sign rule",
                "Extracted closing minor units",
                "Extracted closing raw text",
                "Closing provenance",
                "Source endpoint",
                "Endpoint raw text",
                "Endpoint provenance",
                "Captured authored review JSON",
            ),
            [
                (
                    value.value_id,
                    proof.source_sha256,
                    proof.evidence_sha256,
                    proof.definition_sha256,
                    proof.parser_id,
                    proof.definition_revision,
                    proof.source_identity,
                    proof.source_identity_form,
                    proof.financial_kind,
                    proof.tracking_scope,
                    proof.sign_rule,
                    str(proof.closing.minor),
                    proof.closing.raw,
                    proof.closing.provenance,
                    proof.value_date.iso,
                    proof.value_date.raw,
                    proof.value_date.provenance,
                    proof.review_annotation.model_dump_json() if proof.review_annotation else None,
                )
                for value in source_values
                for proof in (value.source_proof,)
            ],
        )
        _sheet(
            workbook,
            "Value Source Locations",
            ("Value ID", "Source field", "Location ordinal", "Exact location JSON"),
            [
                (
                    value.value_id,
                    field,
                    ordinal,
                    location.model_dump_json(),
                )
                for value in source_values
                for field, locations in (
                    ("closing", value.source_proof.closing.locations),
                    ("endpoint", value.source_proof.value_date.locations),
                )
                for ordinal, location in enumerate(locations, start=1)
            ],
        )
        names.extend(("Source Value Proofs", "Value Source Locations"))
    return names


def write_reviewed_journal_workbook(report: ReviewedJournalExport, destination_path: Path) -> Path:
    """Write active reviewed entry rows, exact posting lines, provenance and scope."""
    destination = destination_path.expanduser()
    if destination.suffix.lower() != ".xlsx":
        raise ReportExportError("report_destination_requires_xlsx")
    if destination.is_symlink():
        raise ReportExportError("report_destination_exists")
    destination = destination.resolve()
    if destination.exists():
        raise ReportExportError("report_destination_exists")
    if not destination.parent.is_dir():
        raise ReportExportError("report_destination_parent_missing")
    accounts = {account.id: account for account in report.accounts}
    source_rows = dict(report.source_rows)
    derived_sources = {
        link.derived_observation_id: link
        for link in report.loan_source_links
        if link.derived_observation_id is not None
    }
    entries = []
    postings = []
    allocations = []
    derived_allocations = []
    for entry in report.entries:
        expense_minor = sum(
            posting.amount_minor for posting in entry.postings if accounts[posting.account_id].kind == "expense"
        )
        income_minor = -sum(
            posting.amount_minor for posting in entry.postings if accounts[posting.account_id].kind == "income"
        )
        entries.append(
            (
                entry.key,
                entry.event_id,
                entry.posting_date.isoformat(),
                entry.description,
                entry.origin,
                entry.reason,
                expense_minor / 100,
                str(expense_minor),
                income_minor / 100,
                str(income_minor),
            )
        )
        for line_number, posting in enumerate(entry.postings, start=1):
            account = accounts[posting.account_id]
            postings.append(
                (
                    entry.key,
                    entry.posting_date.isoformat(),
                    line_number,
                    account.id,
                    account.name,
                    account.kind,
                    posting.amount_minor / 100,
                    str(posting.amount_minor),
                )
            )
            for allocation in posting.allocations:
                if allocation.observation_id in derived_sources:
                    link = derived_sources[allocation.observation_id]
                    if link.entry_key != entry.key or link.account_id != posting.account_id:
                        raise ReportExportError("loan_allocation_source_mismatch")
                    derived_allocations.append(
                        (
                            entry.key,
                            line_number,
                            allocation.observation_id,
                            link.row_id,
                            link.source_sha256,
                            allocation.amount_minor / 100,
                            str(allocation.amount_minor),
                            f"{link.proof.parser_id}: printed principal x {link.proof.sign_factor}",
                        )
                    )
                    continue
                if allocation.observation_id not in source_rows:
                    raise ReportExportError("allocation_source_missing")
                allocations.append(
                    (
                        entry.key,
                        line_number,
                        allocation.observation_id,
                        source_rows[allocation.observation_id],
                        allocation.amount_minor / 100,
                        str(allocation.amount_minor),
                    )
                )
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(
        workbook,
        "Scope",
        ("Field", "Value"),
        [
            ("Report", "ParseTrail 2.0 active reviewed journal"),
            ("From inclusive", report.start.isoformat()),
            ("To exclusive", report.end.isoformat()),
            ("Active reviewed entries", len(report.entries)),
            ("Posting lines", len(postings)),
            ("Unreviewed journal entries excluded in period", report.excluded_unreviewed_entries),
            ("Opening entries excluded in period", report.excluded_opening_entries),
            (
                "Activity source rows without full allocation, whole profile",
                report.activity_rows_without_full_allocation,
            ),
            ("Sources with incomplete activity, whole profile", report.sources_with_incomplete_activity),
            ("Corrections to original entries posted in period", len(report.corrections)),
            ("Retirements of original entries posted in period", len(report.retirements)),
            ("Transfer corrections affecting original events in period", len(report.transfer_actions)),
            ("Source interpretation withdrawals/repairs in period", len(report.source_actions)),
            ("Posted repeat repairs in period", len(report.repeat_repairs)),
            ("Correction presentation", "Superseded entries and reversals are omitted from active rows."),
            (
                "Retirement presentation",
                "Retired originals and their inverses are omitted; links retain history without a replacement.",
            ),
            ("Precision", "USD columns are for display; minor-unit text columns preserve exact values."),
            ("Coverage", "Source rows can overlap; these rows do not prove complete transaction coverage."),
            ("Balances", "This export does not establish reconciled balances or certified cash flow."),
            ("Openings", "Opening positions are not transaction rows and are excluded from this workbook."),
        ],
    )
    _sheet(
        workbook,
        "Reviewed Entries",
        (
            "Entry key",
            "Event ID",
            "Posting date",
            "Description",
            "Origin",
            "Review reason",
            "Net expense USD",
            "Net expense minor units",
            "Net income USD",
            "Net income minor units",
        ),
        entries,
    )
    _sheet(
        workbook,
        "Posting Lines",
        (
            "Entry key",
            "Posting date",
            "Line",
            "Account ID",
            "Account name",
            "Account kind",
            "Signed USD",
            "Signed minor units",
        ),
        postings,
    )
    _sheet(
        workbook,
        "Source Allocations",
        ("Entry key", "Posting line", "Source row ID", "Source SHA-256", "Signed USD", "Signed minor units"),
        allocations,
    )
    _sheet(
        workbook,
        "Corrections",
        ("Original key", "Reversal key", "Replacement key", "Reason", "Created at"),
        [
            (
                correction.original_key,
                correction.reversal_key,
                correction.replacement_key,
                correction.reason,
                correction.created_at,
            )
            for correction in report.corrections
        ],
    )
    _sheet(
        workbook,
        "Retirements",
        ("Original key", "Reversal key", "Reason", "Created at"),
        [(row.original_key, row.reversal_key, row.reason, row.created_at) for row in report.retirements],
    )
    _sheet(
        workbook,
        "Transfer Reviews",
        ("Original event", "Action", "Replacement event", "Reason", "Created at"),
        [
            (row.event_id, row.action, row.replacement_event_id, row.reason, row.created_at)
            for row in report.transfer_actions
        ],
    )
    _sheet(
        workbook,
        "Source Reviews",
        ("Action", "Source row", "Original entry", "Inverse entry", "Kind", "Reason", "Created at"),
        [
            (row.action_id, row.row_id, row.entry_key, row.reversal_key, row.kind, row.reason, row.created_at)
            for row in report.source_actions
        ],
    )
    _sheet(
        workbook,
        "Repeat Repairs",
        (
            "Repair",
            "Retained row",
            "Supporting row",
            "Retained entry",
            "Retired entry",
            "Inverse",
            "Occurrence review",
            "Reason",
            "Created at",
        ),
        [
            (
                row.repair_id,
                row.retained_row_id,
                row.redundant_row_id,
                row.retained_entry_key,
                row.redundant_entry_key,
                row.reversal_key,
                row.review_revision_id,
                row.reason,
                row.created_at,
            )
            for row in report.repeat_repairs
        ],
    )
    expected_sheets = [
        "Scope",
        "Reviewed Entries",
        "Posting Lines",
        "Source Allocations",
        "Corrections",
        "Retirements",
        "Transfer Reviews",
        "Source Reviews",
        "Repeat Repairs",
    ]
    if report.settlements:
        workbook["Scope"].append(
            (
                "Settlement metadata",
                "Full reviewed manifests whose original entry dates intersect this period; may include retired events "
                "and out-of-period legs. Metadata adds no money to active reviewed rows. Remaining clearing is after "
                "all reviewed receipts, not a period-end clearing position.",
            )
        )
        _sheet(
            workbook,
            "Settlement Events",
            (
                "Event ID",
                "Status",
                "Kind",
                "Outgoing row",
                "Sent principal minor units",
                "Received gross minor units",
                "Remaining clearing minor units",
                "Review reason",
                "Created at",
                "Version",
                "Replacement event",
                "Action reason",
                "Original entry references",
                "Expected owned destination ID",
                "Expected reviewed name",
                "Expected kind",
                "Expected tracking scope",
                "Ownership acknowledged",
                "Ownership review reference",
                "Expectation qualification",
            ),
            [
                (
                    event.event_id,
                    event.status,
                    event.settlement.kind,
                    event.settlement.outgoing.row_id,
                    str(event.settlement.sent_principal_minor),
                    str(event.settlement.received_gross_minor),
                    str(event.settlement.residual_minor),
                    event.reason,
                    event.created_at,
                    event.version,
                    event.replacement_event_id,
                    event.action_reason,
                    "\n".join(event.entry_keys),
                    event.settlement.expected_destination.account_id if event.settlement.expected_destination else None,
                    event.settlement.expected_destination.account_name
                    if event.settlement.expected_destination
                    else None,
                    event.settlement.expected_destination.account_kind
                    if event.settlement.expected_destination
                    else None,
                    event.settlement.expected_destination.tracking_scope
                    if event.settlement.expected_destination
                    else None,
                    event.settlement.expected_destination.ownership_confirmed
                    if event.settlement.expected_destination
                    else None,
                    event.settlement.expected_destination.reference if event.settlement.expected_destination else None,
                    "Authored ownership expectation; no receipt evidence or destination posting"
                    if event.settlement.expected_destination
                    else None,
                )
                for event in report.settlements
            ],
        )
        _sheet(
            workbook,
            "Settlement Legs",
            (
                "Event ID",
                "Status",
                "Leg",
                "Source row",
                "Source SHA-256",
                "Account ID",
                "Account name",
                "Account kind",
                "Accounting date",
                "Date provenance",
                "Raw signed source minor units",
                "Reviewed transfer portion minor units",
                "Tracking scope",
                "Source description known",
            ),
            [
                (
                    event.event_id,
                    event.status,
                    "outgoing" if index == 0 else "receipt",
                    movement.row_id,
                    movement.source_sha256,
                    movement.account_id,
                    movement.account_name,
                    movement.account_kind,
                    movement.posting_date,
                    movement.date_provenance,
                    str(movement.amount_minor),
                    str(
                        event.settlement.sent_principal_minor
                        if index == 0
                        else movement.amount_minor
                        + sum(fee.amount_minor for fee in event.settlement.fees_for(movement.row_id))
                    ),
                    movement.tracking_scope,
                    movement.source_description_known,
                )
                for event in report.settlements
                for index, movement in enumerate(event.settlement.movements)
            ],
        )
        _sheet(
            workbook,
            "Settlement Fees",
            (
                "Event ID",
                "Status",
                "Source row",
                "Accounting date",
                "Expense category ID",
                "Category name",
                "Exact fee minor units",
                "Source/review reference",
                "Allocation reason",
            ),
            [
                (
                    event.event_id,
                    event.status,
                    fee.source_row_id,
                    next(
                        movement.posting_date
                        for movement in event.settlement.movements
                        if movement.row_id == fee.source_row_id
                    ),
                    fee.category_id,
                    accounts[fee.category_id].name,
                    str(fee.amount_minor),
                    fee.reference,
                    fee.reason,
                )
                for event in report.settlements
                for fee in event.settlement.fees
            ],
        )
        expected_sheets.extend(("Settlement Events", "Settlement Legs", "Settlement Fees"))
    expected_sheets.extend(_loan_metadata(workbook, report, derived_allocations))
    expected_sheets.extend(_value_metadata(workbook, report, accounts))
    descriptor, partial_name = tempfile.mkstemp(
        prefix=f".{destination.name}.partial.", suffix=".xlsx", dir=destination.parent
    )
    os.close(descriptor)
    partial = Path(partial_name)
    try:
        workbook.save(partial)
        verified = load_workbook(partial, read_only=True, data_only=False)
        try:
            if verified.sheetnames != expected_sheets:
                raise ReportExportError("report_workbook_integrity")
        finally:
            verified.close()
        try:
            os.link(partial, destination)
        except FileExistsError:
            raise ReportExportError("report_destination_exists") from None
        return destination
    finally:
        partial.unlink(missing_ok=True)
