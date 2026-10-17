"""Bounded, read-only XLSX decoder for the isolated table evidence pilot."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable, Iterator

import openpyxl

from parsetrail.core.v2_parser_engine import (
    MAX_CELL_CHARS,
    MAX_CELLS,
    MAX_COLUMNS,
    MAX_ROWS,
    MAX_SOURCE_BYTES,
    EngineError,
    Evidence,
    load_definition,
    parse_table,
)

MAX_ARCHIVE_MEMBERS = 256
MAX_EXPANDED_BYTES = 16 * 1024 * 1024
MAX_MEMBER_BYTES = 8 * 1024 * 1024
MAX_SHEETS = 32
READ_CHUNK = 64 * 1024


class XlsxError(ValueError):
    """Redacted workbook rejection without source cell contents."""

    def __init__(self, code: str, row: int | None = None):
        self.code = code
        self.row = row
        super().__init__(code)


def _inspect_archive(source: bytes, cancelled: Callable[[], bool] | None) -> None:
    """Read every member through explicit expansion limits before XML decoding."""
    try:
        with zipfile.ZipFile(io.BytesIO(source)) as archive:
            members = archive.infolist()
            if not members or len(members) > MAX_ARCHIVE_MEMBERS:
                raise XlsxError("archive_members")
            names: set[str] = set()
            expanded = 0
            for member in members:
                if cancelled is not None and cancelled():
                    raise XlsxError("cancelled")
                name = member.filename
                if (
                    not name
                    or name in names
                    or name.startswith("/")
                    or "\\" in name
                    or ":" in name
                    or any(part in {"", ".", ".."} for part in name.rstrip("/").split("/"))
                    or member.flag_bits & 1
                    or member.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
                ):
                    raise XlsxError("archive_member")
                names.add(name)
                if member.is_dir():
                    continue
                if member.file_size > MAX_MEMBER_BYTES or expanded + member.file_size > MAX_EXPANDED_BYTES:
                    raise XlsxError("expanded_size")
                actual = 0
                with archive.open(member) as stream:
                    while chunk := stream.read(READ_CHUNK):
                        if cancelled is not None and cancelled():
                            raise XlsxError("cancelled")
                        actual += len(chunk)
                        if actual > MAX_MEMBER_BYTES or expanded + actual > MAX_EXPANDED_BYTES:
                            raise XlsxError("expanded_size")
                if actual != member.file_size:
                    raise XlsxError("archive_member_size")
                expanded += actual
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise XlsxError("xlsx_structure")
    except (zipfile.BadZipFile, RuntimeError, OSError, EOFError):
        raise XlsxError("xlsx_structure") from None


def _rows(sheet, cancelled: Callable[[], bool] | None) -> Iterator[tuple[str | None, ...]]:
    for row_number, cells in enumerate(sheet.iter_rows(), 1):
        if cancelled is not None and cancelled():
            raise XlsxError("cancelled", row_number)
        if row_number > MAX_ROWS:
            raise XlsxError("row_limit", row_number)
        values = []
        for cell in cells:
            if cell.data_type == "f":
                raise XlsxError("formula_cell", row_number)
            value = cell.value
            if value is not None and (not isinstance(value, str) or len(value) > MAX_CELL_CHARS):
                raise XlsxError("cell_type_or_size", row_number)
            values.append(value)
        yield tuple(values)


def parse_xlsx(
    definition_bytes: bytes,
    source_bytes: bytes,
    *,
    sheet: str,
    cancelled: Callable[[], bool] | None = None,
) -> Evidence:
    """Decode a selected string-cell sheet and feed the finite table pilot."""
    if not source_bytes or len(source_bytes) > MAX_SOURCE_BYTES:
        raise XlsxError("source_size")
    definition, _ = load_definition(definition_bytes)
    if definition.suffix != ".xlsx" or not sheet or len(sheet) > 31:
        raise XlsxError("source_kind")
    _inspect_archive(source_bytes, cancelled)
    workbook = None
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(source_bytes), read_only=True, data_only=False, keep_links=False)
        if len(workbook.sheetnames) > MAX_SHEETS or sheet not in workbook.sheetnames:
            raise XlsxError("missing_or_excess_sheets")
        selected = workbook[sheet]
        if (
            selected.max_row is None
            or selected.max_column is None
            or selected.max_row > MAX_ROWS
            or selected.max_column > MAX_COLUMNS
            or selected.max_row * selected.max_column > MAX_CELLS
        ):
            raise XlsxError("sheet_dimensions")
        return parse_table(
            definition_bytes,
            source_bytes,
            _rows(selected, cancelled),
            suffix=".xlsx",
            sheet=sheet,
            extraction_adapter=f"openpyxl-{openpyxl.__version__}-read-only",
            cancelled=cancelled,
        )
    except (EngineError, XlsxError):
        raise
    except Exception:
        raise XlsxError("xlsx_structure") from None
    finally:
        if workbook is not None:
            workbook.close()
